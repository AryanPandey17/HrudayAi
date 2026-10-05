"""Load the raw Z-Alizadeh Sani extension workbook into a DataFrame."""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from ml.config import LEAKAGE_COLUMNS, RAW_DATA_PATH, TARGETS


def load_raw(path: Path = RAW_DATA_PATH) -> pd.DataFrame:
    """Read the first worksheet with openpyxl; the first row is the header."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = [str(name).strip() for name in next(rows)]
        frame = pd.DataFrame(list(rows), columns=header)
    finally:
        workbook.close()
    _require_columns(frame, LEAKAGE_COLUMNS)
    return frame


def feature_columns(frame: pd.DataFrame) -> list[str]:
    """Columns allowed as model inputs: everything except the leakage columns."""
    return [column for column in frame.columns if column not in LEAKAGE_COLUMNS]


def target_series(frame: pd.DataFrame, target: str) -> pd.Series:
    """Binary 0/1 labels for a target; raises on any label outside the target's two classes."""
    spec = TARGETS[target]
    values = frame[spec.column]
    unexpected = set(values.unique()) - {spec.positive_label, spec.negative_label}
    if unexpected:
        raise ValueError(f"Unexpected labels in '{spec.column}': {sorted(map(str, unexpected))}")
    return (values == spec.positive_label).astype(int).rename(target)


def _require_columns(frame: pd.DataFrame, required: frozenset[str]) -> None:
    missing = sorted(required - set(frame.columns))
    if missing:
        raise ValueError(f"Dataset is missing required columns: {missing}")
